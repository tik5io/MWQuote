import os
import tempfile
import unittest

from domain import offers
from domain.phase import PhaseKind
from infrastructure.persistence import PersistenceService
from tests.test_phases import _project


class RevisionTest(unittest.TestCase):
    def test_next_revision_follows_the_alphabet(self):
        history = []
        self.assertEqual(offers.next_revision(history, "OD1"), "A")
        history.append({"devis_ref": "OD1", "revision": "A"})
        self.assertEqual(offers.next_revision(history, "OD1"), "B")
        history.append({"devis_ref": "OD1", "revision": "B"})
        history.append({"devis_ref": "OD2", "revision": "A"})
        self.assertEqual(offers.next_revision(history, "OD1"), "C")
        self.assertEqual(offers.next_revision(history, "OD2"), "B")

    def test_legacy_entry_counts_as_revision_a(self):
        history = [{"devis_ref": "OD1"}]          # ancienne entrée : pas de champ revision
        self.assertEqual(offers.revision_of(history[0]), "A")
        self.assertEqual(offers.next_revision(history, "OD1"), "B")

    def test_revision_names_roll_over_after_z(self):
        history = [{"devis_ref": "OD1", "revision": "Z"}]
        self.assertEqual(offers.next_revision(history, "OD1"), "AA")

    def test_group_offers_orders_revisions(self):
        history = [
            {"devis_ref": "OD1", "version_index": 1, "revision": "B"},
            {"devis_ref": "OD2", "version_index": 2, "revision": "A"},
            {"devis_ref": "OD1", "version_index": 1, "revision": "A"},
        ]
        groups = offers.group_offers(history)
        self.assertEqual([(g["devis_ref"], g["version_index"]) for g in groups], [("OD1", 1), ("OD2", 2)])
        self.assertEqual([offers.revision_of(e) for e in groups[0]["entries"]], ["A", "B"])


class FreezeTest(unittest.TestCase):
    def test_any_offer_freezes_its_version_legacy_ones_included(self):
        legacy = {"devis_ref": "OLD", "version_index": 1}            # offre antérieure aux indices
        self.assertTrue(offers.is_legacy(legacy))
        self.assertFalse(offers.is_version_frozen([], 1))
        self.assertTrue(offers.is_version_frozen([legacy], 1))       # elle fige quand même sa version
        self.assertFalse(offers.is_version_frozen([legacy], 2))
        self.assertEqual(offers.frozen_by([legacy], 1), "OLD")
        tracked = {"devis_ref": "OD1", "version_index": 1, "revision": "A"}
        self.assertFalse(offers.is_legacy(tracked))
        self.assertTrue(offers.is_version_frozen([tracked], 1))

    def test_legacy_offers_of_a_version_are_grouped_under_it(self):
        history = [{"devis_ref": "OLD-1", "version_index": 1}, {"devis_ref": "OLD-2", "version_index": 1},
                   {"devis_ref": "OD3", "version_index": 1, "revision": "A"}]
        groups = offers.group_offers(history)
        self.assertEqual([g["version_index"] for g in groups], [1, 1, 1])
        self.assertEqual([g["devis_ref"] for g in groups], ["OLD-1", "OLD-2", "OD3"])

    def test_legacy_project_is_locked_after_load(self):
        p = _project()
        p.export_history.append({"devis_ref": "OLD", "version_index": 1, "date": "01/01/2026"})   # sans revision
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "p.mwq")
            PersistenceService.save_project(p, path)
            q = PersistenceService.load_project(path)
        self.assertTrue(q.is_version_frozen())
        self.assertEqual(q.version_frozen_by(), "OLD")

    def test_new_version_of_a_frozen_one_is_free(self):
        p = _project()
        p.export_history.append({"devis_ref": "OD1", "version_index": 1, "revision": "A"})
        self.assertTrue(p.is_version_frozen())
        v2 = p.add_version("V2")
        self.assertFalse(p.is_version_frozen(v2.version_index))
        self.assertEqual(p.version_frozen_by(1), "OD1")


class SnapshotTest(unittest.TestCase):
    def test_lot_snapshot_roundtrip(self):
        p = _project()
        version = p.current_version
        snap = offers.snapshot_coefficients(version)
        self.assertEqual(snap, {"mode": "lot", "quantities": [10, 100], "rates": {"10": 1.2, "100": 1.1}})
        version.phase.volume_margin_rates = {10: 1.5, 100: 1.4}
        version.phase.sale_quantities = [10, 100, 500]
        self.assertFalse(offers.same_coefficients(snap, offers.snapshot_coefficients(version)))
        self.assertTrue(offers.apply_snapshot(version, snap))
        self.assertEqual(version.sale_quantities, [10, 100])
        self.assertEqual(version.volume_margin_rates, {10: 1.2, 100: 1.1})
        self.assertTrue(offers.same_coefficients(snap, offers.snapshot_coefficients(version)))

    def test_decompose_snapshot_roundtrip(self):
        p = _project()
        v = p.add_version_variant("Déc", PhaseKind.SERIE, annual_volume=50000)
        v.serie_data.global_commercial_margin = 0.30
        snap = offers.snapshot_coefficients(v)
        v.serie_data.annual_volume = 70000
        v.serie_data.global_commercial_margin = 0.10
        self.assertTrue(offers.apply_snapshot(v, snap))
        self.assertEqual((v.serie_data.annual_volume, v.serie_data.global_commercial_margin), (50000, 0.30))

    def test_snapshot_of_the_wrong_mode_is_refused(self):
        p = _project()
        lot_snap = offers.snapshot_coefficients(p.current_version)
        v = p.add_version_variant("Déc", PhaseKind.SERIE, annual_volume=1000)
        self.assertFalse(offers.apply_snapshot(v, lot_snap))
        self.assertFalse(offers.apply_snapshot(p.current_version, None))


class LatestOfferRefTest(unittest.TestCase):
    def test_latest_offer_ref_per_version(self):
        h = [{"devis_ref": "A1", "version_index": 1}, {"devis_ref": "B2", "version_index": 2},
             {"devis_ref": "A1", "version_index": 1, "revision": "B"}]
        self.assertEqual(offers.latest_offer_ref(h, 1), "A1")
        self.assertEqual(offers.latest_offer_ref(h, 2), "B2")
        self.assertIsNone(offers.latest_offer_ref(h, 3))
        self.assertIsNone(offers.latest_offer_ref([], 1))


class ExportRevisionTest(unittest.TestCase):
    """Génération réelle d'un XLSX avec le gabarit de l'application."""

    def setUp(self):
        from core.app_icon import get_template_path
        from infrastructure.export_service import ExportService
        self.template = str(get_template_path())
        self.service = ExportService()          # sans base : numérotation de repli
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _export(self, project, name, **kw):
        path = os.path.join(self.tmp.name, name)
        self.service.export_excel(project, self.template, path, **kw)
        return path

    def test_usd_export_converts_prices_and_records_rate(self):
        from openpyxl import load_workbook
        from infrastructure.fx_rate import FxRate
        p = _project()
        saved = (FxRate.rate, FxRate.currency)
        FxRate.rate = 1.25
        try:
            path = self._export(p, "usd.xlsx", devis_ref="OD261023_002", currency="USD")
        finally:
            FxRate.rate, FxRate.currency = saved
        entry = p.export_history[-1]
        self.assertEqual((entry["currency"], entry["fx_rate"]), ("USD", 1.25))
        prices = [c.value for row in load_workbook(path).active.iter_rows() for c in row
                  if c.number_format.startswith("$") and isinstance(c.value, (int, float))]
        self.assertTrue(prices)
        self.assertAlmostEqual(prices[0], p.total_price(sorted(p.sale_quantities)[0]) * 1.25, places=6)

    def test_first_export_is_revision_a_with_snapshot_then_b(self):
        p = _project()
        self._export(p, "a.xlsx", devis_ref="OD261023_001")
        a = p.export_history[-1]
        self.assertEqual((a["devis_ref"], a["revision"]), ("OD261023_001", "A"))
        self.assertEqual(a["coeffs"]["rates"], {"10": 1.2, "100": 1.1})

        p.volume_margin_rates[10] = 1.35                      # ajustement de coefficient
        self._export(p, "b.xlsx", devis_ref="OD261023_001")
        b = p.export_history[-1]
        self.assertEqual((b["devis_ref"], b["revision"]), ("OD261023_001", "B"))
        self.assertEqual(b["coeffs"]["rates"]["10"], 1.35)
        self.assertEqual(a["coeffs"]["rates"]["10"], 1.2)     # l'indice A n'a pas bougé
        self.assertTrue(p.is_version_frozen())

    def test_draft_adds_nothing_and_freezes_nothing(self):
        p = _project()
        self._export(p, "d.xlsx", draft=True)
        self.assertEqual(p.export_history, [])
        self.assertFalse(p.is_version_frozen())

    def test_default_filename_and_display_ref(self):
        p = _project()
        self.assertEqual(self.service.display_ref("OD1_001", "A"), "OD1_001")
        self.assertEqual(self.service.display_ref("OD1_001", "B"), "OD1_001 ind. B")
        name_a = self.service.get_default_filename(p, devis_ref="OD1_001")
        name_b = self.service.get_default_filename(p, devis_ref="OD1_001", revision="B")
        self.assertTrue(name_a.startswith("OD1_001-REF"))
        self.assertTrue(name_b.startswith("OD1_001-indB-REF"))

    def test_revision_and_snapshot_survive_save_and_load(self):
        p = _project()
        self._export(p, "a.xlsx", devis_ref="OD261023_002")
        path = os.path.join(self.tmp.name, "p.mwq")
        PersistenceService.save_project(p, path)
        q = PersistenceService.load_project(path)
        entry = q.export_history[0]
        self.assertEqual((entry["revision"], entry["coeffs"]["quantities"]), ("A", [10, 100]))
        self.assertTrue(q.is_version_frozen())


class OfferWorkbookNameTest(unittest.TestCase):
    """Excel refuse deux classeurs de même nom ouverts ensemble : nom = référence d'offre (unique)."""

    def _entry(self, ref, revision=None, filename="Offre.xlsx", payload=b"PK-data"):
        import base64
        e = {"devis_ref": ref, "xlsx_filename": filename, "xlsx_data_b64": base64.b64encode(payload).decode()}
        if revision:
            e["revision"] = revision
        return e

    def test_name_comes_from_offer_ref_and_revision(self):
        from infrastructure.offer_files import workbook_name
        self.assertEqual(workbook_name(self._entry("OD261016_004", "A")), "OD261016_004.xlsx")
        self.assertEqual(workbook_name(self._entry("OD261016_004", "B")), "OD261016_004-indB.xlsx")
        self.assertEqual(workbook_name(self._entry("OD261016_004")), "OD261016_004.xlsx")      # ancienne offre
        self.assertEqual(workbook_name(self._entry("OD 26/10:4")), "OD_26_10_4.xlsx")          # caractères interdits

    def test_two_offers_with_the_same_user_filename_get_distinct_workbooks(self):
        from infrastructure.offer_files import offer_workbook_path
        with tempfile.TemporaryDirectory() as tmp:
            a = offer_workbook_path(self._entry("OD1_001", "A", filename="Devis.xlsx", payload=b"A"), tmp)
            b = offer_workbook_path(self._entry("OD1_002", "A", filename="Devis.xlsx", payload=b"B"), tmp)
            c = offer_workbook_path(self._entry("OD1_001", "B", filename="Devis.xlsx", payload=b"C"), tmp)
            self.assertEqual(len({os.path.basename(x) for x in (a, b, c)}), 3)
            self.assertEqual([open(x, "rb").read() for x in (a, b, c)], [b"A", b"B", b"C"])

    def test_workbook_already_open_in_excel_is_reused_or_suffixed(self):
        import builtins
        from unittest import mock
        from infrastructure import offer_files
        real_open = builtins.open
        with tempfile.TemporaryDirectory() as tmp:
            entry = self._entry("OD1_001", "A", payload=b"v1")
            path = offer_files.offer_workbook_path(entry, tmp)

            def locked_open(file, mode="r", *args, **kwargs):      # Excel tient le fichier : écriture refusée
                if os.fspath(file) == path and "w" in mode:
                    raise PermissionError(13, "fichier utilisé par un autre processus")
                return real_open(file, mode, *args, **kwargs)

            with mock.patch.object(offer_files, "open", locked_open, create=True):
                self.assertEqual(offer_files.offer_workbook_path(entry, tmp), path)          # même contenu : réutilisé
                changed = self._entry("OD1_001", "A", payload=b"v2")
                alt = offer_files.offer_workbook_path(changed, tmp)                          # contenu différent : variante
            self.assertNotEqual(alt, path)
            self.assertTrue(os.path.basename(alt).startswith("OD1_001~"))
            self.assertEqual(real_open(alt, "rb").read(), b"v2")

    def test_entry_without_file_is_refused(self):
        from infrastructure.offer_files import offer_workbook_path
        with self.assertRaises(ValueError):
            offer_workbook_path({"devis_ref": "OD1"}, tempfile.gettempdir())


if __name__ == "__main__":
    unittest.main()
