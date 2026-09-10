"""
Tests pour l'export de synthèse multi-offres (ExportService.export_multi_synthesis).
Vérifie la structure du classeur : feuille Synthèse + une feuille par offre,
avec les prix de vente unitaires par quantité.
"""

import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from domain.project import Project
from domain.operation import Operation
from domain.cost import CostItem, CostType, PricingStructure, PricingType
from infrastructure.export_service import ExportService


def _make_project(reference, client, unit_price, quantities):
    """Projet minimal avec une opération et un coût matière par unité."""
    cost = CostItem(
        name="Matière",
        cost_type=CostType.MATERIAL,
        pricing=PricingStructure(PricingType.PER_UNIT, unit_price=unit_price),
        quantity_per_piece=1.0,
        margin_rate=0.0,
    )
    op = Operation(code="OP10", label="Usinage", typology="Interne", costs={"Matière": cost})
    return Project(
        name=f"Projet {reference}",
        reference=reference,
        client=client,
        operations=[op],
        sale_quantities=list(quantities),
    )


def test_synthesis_workbook_structure():
    p1 = _make_project("REF-A", "Client Alpha", unit_price=2.0, quantities=[10, 100])
    p2 = _make_project("REF-B", "Client Beta", unit_price=5.0, quantities=[100, 500])

    export = ExportService(db=None)

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "synthese.xlsx"
        export.export_multi_synthesis([p1, p2], str(out))
        assert out.exists()

        wb = load_workbook(str(out))
        # Feuille synthèse + une feuille par offre
        assert "Synthèse" in wb.sheetnames
        assert "REF-A" in wb.sheetnames
        assert "REF-B" in wb.sheetnames

        # Union des quantités -> colonnes 10, 100, 500 (à partir de la colonne D)
        synth = wb["Synthèse"]
        header = [synth.cell(row=4, column=c).value for c in range(1, 7)]
        assert header[0] == "Référence"
        assert header[1] == "Client"
        assert header[2] == "Mode"

        # La feuille détail contient une ligne TOTAL / pièce égale à total_price(qty).
        detail = wb["REF-A"]
        total_cells = []
        for row in detail.iter_rows():
            for cell in row:
                if cell.value == "TOTAL / pièce":
                    total_cells.append(cell.row)
        assert total_cells, "Ligne TOTAL / pièce manquante dans la feuille détail"

        total_row = total_cells[0]
        # colonnes C, D pour quantités 10 et 100
        val_q10 = detail.cell(row=total_row, column=3).value
        val_q100 = detail.cell(row=total_row, column=4).value
        assert val_q10 == pytest.approx(p1.total_price(10), rel=1e-4)
        assert val_q100 == pytest.approx(p1.total_price(100), rel=1e-4)


def test_synthesis_requires_projects():
    export = ExportService(db=None)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "empty.xlsx"
        with pytest.raises(Exception):
            export.export_multi_synthesis([], str(out))


def test_synthesis_sheet_names_deduplicated():
    # Deux offres avec la même référence -> noms de feuilles uniques.
    p1 = _make_project("SAME", "C1", unit_price=1.0, quantities=[10])
    p2 = _make_project("SAME", "C2", unit_price=2.0, quantities=[10])

    export = ExportService(db=None)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "dup.xlsx"
        export.export_multi_synthesis([p1, p2], str(out))
        wb = load_workbook(str(out))
        # "Synthèse" + deux feuilles distinctes pour les deux offres homonymes
        offer_sheets = [s for s in wb.sheetnames if s != "Synthèse"]
        assert len(offer_sheets) == 2
        assert len(set(offer_sheets)) == 2
